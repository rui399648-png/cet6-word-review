# -*- coding: utf-8 -*-
"""独立校验导出的打卡 xlsx：用法 python verify_export.py <导出文件> <原始文件>
结构断言：行数一致、表头一致、A-E 数据列一致、轮次列只含 ☑/□、掌握程度 ∈ 四档。
输出以 OK / FAIL 开头的行；失败时 exit 1。"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

MASTERY = {"陌生", "模糊", "熟悉", "掌握"}
MARKS = {"☑", "□"}

def load_rows(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    return wb.sheetnames[0], [list(r) for r in ws.iter_rows(values_only=True)]

def main():
    out_path, orig_path = sys.argv[1], sys.argv[2]
    sheet_out, rows = load_rows(out_path)
    sheet_orig, orig = load_rows(orig_path)
    fails = []
    if sheet_out != sheet_orig:
        fails.append(f"sheet name {sheet_out!r} != {sheet_orig!r}")
    if len(rows) != len(orig):
        fails.append(f"row count {len(rows)} != {len(orig)}")
    if [str(x) for x in rows[0]] != [str(x) for x in orig[0]]:
        fails.append("header mismatch")
    for i in range(1, min(len(rows), len(orig))):
        for c in range(0, 5):  # 序号/单词/词性/释义/来源 原样
            if str(rows[i][c]) != str(orig[i][c]):
                fails.append(f"row{i+1} col{c+1} data changed")
        for c in range(5, 10):  # 轮次列
            v = "" if rows[i][c] is None else str(rows[i][c])
            if v not in MARKS:
                fails.append(f"row{i+1} col{c+1} bad mark {v!r}")
        m = "" if rows[i][10] is None else str(rows[i][10])
        if m not in MASTERY:
            fails.append(f"row{i+1} mastery bad {m!r}")
    if fails:
        print("FAIL " + "; ".join(fails[:10]))
        sys.exit(1)
    print(f"OK verify_export rows={len(rows)} marks_and_mastery_valid")

if __name__ == "__main__":
    main()
